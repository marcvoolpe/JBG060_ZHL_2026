"""Fill the sprint-review template with the cropland-audit problem, background and first results.

Reads course_content/Template_sprint_review.pptx (never modified) and writes
course_content/Sprint_review_cropland_draft.pptx. Slides 1-3, 7 and 10 are kept as
the team left them. The SLE slides (8-9 plus three new ones) were added at the team's request:
they report project facts and point to lecture concepts; the essay itself stays the team's own work.

Usage (repo root, global Python 3.13):
    python deliverables/sprint_review/build_sprint_review_deck.py
"""

from __future__ import annotations

import copy
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
import rasterio
from matplotlib.colors import ListedColormap
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.enum.dml import MSO_THEME_COLOR
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[2]
TEMPLATE = ROOT / "course_content" / "Template_sprint_review.pptx"
OUT = ROOT / "course_content" / "Sprint_review_cropland_draft.pptx"
DEMO = ROOT / "data" / "cropland" / "demo"
FIG = ROOT / "deliverables" / "sprint_review"

RED = RGBColor(0xC0, 0x00, 0x00)
INK = RGBColor(0x1F, 0x1F, 0x1F)
MUTED = RGBColor(0x5F, 0x63, 0x68)
TINT = RGBColor(0xE9, 0xF5, 0xF6)      # light teal card
TEAL = RGBColor(0x2E, 0x86, 0x8F)      # dark teal for text on tint
GOLD = RGBColor(0xD4, 0xA0, 0x17)      # "map says crop" colour, as in the figures
GREY = RGBColor(0x9A, 0x9A, 0x9A)


# --- helpers -----------------------------------------------------------------------

def set_title(slide, template_title_txbody, line1: str, line2: str) -> None:
    """Copy the template's two-tone title (black + red) and replace the words."""
    title = slide.shapes.title
    old = title._element.txBody
    new = copy.deepcopy(template_title_txbody)
    old.getparent().replace(old, new)
    runs = title.text_frame.paragraphs[0].runs
    runs[0].text = line1.rstrip()
    runs[1].text = line2
    # always break after the section name, so the red subtitle starts on its own line
    br = runs[0]._r.makeelement("{http://schemas.openxmlformats.org/drawingml/2006/main}br", {})
    runs[0]._r.addnext(br)


def drop_body(slide) -> None:
    for ph in list(slide.placeholders):
        if ph.placeholder_format.idx != 0:
            ph._element.getparent().remove(ph._element)


def text(slide, x, y, w, h, paras, size=14, color=INK, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
         italic=False):
    """paras: list of str or (str, dict) with per-paragraph overrides: size, color, bold, italic, space."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.margin_top = tf.margin_bottom = Inches(0.03)
    for i, p in enumerate(paras):
        s, o = (p, {}) if isinstance(p, str) else p
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = o.get("align", align)
        para.space_after = Pt(o.get("space", 4))
        r = para.add_run()
        r.text = s
        f = r.font
        f.size = Pt(o.get("size", size))
        f.bold = o.get("bold", bold)
        f.italic = o.get("italic", italic)
        f.color.rgb = o.get("color", color)
    return tb


def bullets(slide, x, y, w, h, items, size=14, color=INK, space=6):
    """Real bullets (no literal characters): items are str or (str, level)."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.05)
    for i, it in enumerate(items):
        s, lvl = (it, 0) if isinstance(it, str) else it
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.space_after = Pt(space)
        pPr = para._p.get_or_add_pPr()
        indent = 228600 if lvl == 0 else 457200
        pPr.set("marL", str(indent))
        pPr.set("indent", str(-228600))
        bu = pPr.makeelement("{http://schemas.openxmlformats.org/drawingml/2006/main}buChar", {"char": "•"})
        pPr.append(bu)
        r = para.add_run()
        r.text = s
        r.font.size = Pt(size - 2 * lvl)
        r.font.color.rgb = color if lvl == 0 else MUTED
    return tb


def card(slide, x, y, w, h, fill=TINT):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.adjustments[0] = 0.08
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    s.line.fill.background()
    s.shadow.inherit = False
    return s


def circle_num(slide, x, y, d, label, fill=TEAL):
    s = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    s.line.fill.background()
    s.shadow.inherit = False
    tf = s.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = label
    r.font.size = Pt(14)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    return s


def arrow(slide, x1, y1, x2, y2):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = GREY
    c.line.width = Pt(2)
    ln = c.line._get_or_add_ln()
    tail = ln.makeelement("{http://schemas.openxmlformats.org/drawingml/2006/main}tailEnd", {"type": "triangle"})
    ln.append(tail)
    return c


def source(slide, s):
    text(slide, 0.69, 5.2, 8.62, 0.3, [s], size=9, color=MUTED)


# --- figures -----------------------------------------------------------------------

def demo_figure() -> Path:
    """Large-font version of the demo panels (outside the study counties) for a slide."""
    rgb = rasterio.open(DEMO / "sentinel2_oct_dec_2022_rgb.tif").read().astype(float)
    rgb = np.clip(np.moveaxis(rgb, 0, -1) / 3000.0, 0, 1)
    panels = [("WorldCover 2021", "worldcover.tif", 40), ("Esri 2022", "esri2022.tif", 5),
              ("WorldCereal 2021", "worldcereal.tif", 100), ("Dyn. World 2022", "dw_jjason.tif", 4),
              ("GLAD 2019", "glad2019.tif", 1), ("DE Africa 2019", "deafrica_mask.tif", 1),
              ("GFSAD 2015", "gfsad2015.tif", 2)]
    fig, axes = plt.subplots(2, 4, figsize=(12.8, 6.4), constrained_layout=True)
    axes = axes.ravel()
    axes[0].imshow(rgb)
    axes[0].set_title("Satellite image\n(Sentinel-2, late 2022)", fontsize=15)
    cmap = ListedColormap(["#ece7dc", "#d4a017", "#ffffff"])
    for ax, (name, f, code) in zip(axes[1:], panels):
        a = rasterio.open(DEMO / f).read(1)
        disp = np.where(a == 255, 2, (a == code).astype(int))
        share = 100 * (a[a != 255] == code).mean()
        ax.imshow(disp, cmap=cmap, vmin=0, vmax=2, interpolation="nearest")
        ax.set_title(f"{name}\ncrop = {share:.1f}%", fontsize=15)
    for ax in axes:
        ax.set_xticks([])
        ax.set_yticks([])
    path = FIG / "slide_demo_maps_turalei.png"
    fig.savefig(path, dpi=150)
    plt.close(fig)
    return path


def aweil_shares() -> pd.DataFrame:
    s = pd.read_csv(ROOT / "data" / "cropland" / "agreement" / "county_crop_share.csv")
    a = s[s.aoi == "aweil"].groupby("product").agg(ha=("crop_area_ha_10m", "sum"), valid=("valid_ha", "sum"))
    a["share"] = 100 * a.ha / a.valid
    c = pd.read_csv(ROOT / "data" / "cropland" / "county_cropland_areas.csv")
    aw = c[c.aoi == "aweil"]
    area = (aw.county_area_km2 * 100).sum()
    names = {"dw2022jjason": "Dynamic World 2022", "worldcover2021": "WorldCover 2021", "esri2022": "Esri 2022",
             "worldcereal2021": "WorldCereal 2021", "glad2019": "GLAD 2019", "deafrica2019": "DE Africa 2019",
             "gfsad2015": "GFSAD 2015"}
    rows = [(names[k], a.loc[k, "share"]) for k in names]
    rows.append(("ASAP (course data)", 100 * aw.asap_crop_v04_ha.sum() / area))
    rows = sorted(rows, key=lambda r: r[1])
    rows.append(("CFSAM 2022 (modelled)", 100 * aw.cfsam2022_cereal_harvested_ha_MODELLED.sum() / area))
    return pd.DataFrame(rows, columns=["source", "share"])


# --- slides ------------------------------------------------------------------------

def slide_problem(s):
    bullets(s, 0.69, 1.55, 4.3, 3.5, [
        "Anticipatory action for floods needs to know how much farmland is under water: ZOA's example trigger is "
        "a flood “inundating cropland”.",
        "In South Sudan that number comes from overlaying public flood maps on public cropland maps.",
        "Nobody has checked how accurate those maps are where ZOA works: along the Lol river (Aweil) and in "
        "Bor South.",
    ], size=15, space=10)
    card(s, 5.25, 1.55, 4.06, 3.45)
    text(s, 5.45, 1.68, 3.7, 0.4, ["Research objective"], size=16, bold=True, color=TEAL)
    text(s, 5.45, 2.1, 3.7, 2.1, [
        "Evaluate how accurately public satellite data measure cropland, and cropland flooding, in ZOA's areas, "
        "and which conditions go with flooding of farmland."], size=14)
    text(s, 5.45, 3.85, 3.7, 1.1, [
        ("Five Aweil counties + Bor South", {"size": 11, "color": MUTED}),
        ("Reference season 2022, flood record 2017–2025", {"size": 11, "color": MUTED}),
        ("≈450 blind, double-labelled 30 m points", {"size": 11, "color": MUTED}),
    ], size=11)
    s.notes_slide.notes_text_frame.text = (
        "Trigger wording: Stakeholder Q&A 2, slides 31 and 40 (\"a flood inundating cropland\"). Areas: slide 3. "
        "RO as in the team's final plan of 27 Sep 2026; forecasting is out of scope.")


def slide_paper(s):
    y = 1.6
    boxes = [("Flood map", "satellite (MODIS)"), ("Cropland map", "land use"), ("Flooded cropland", "hectares"),
             ("Crop & calorie loss", "+ agricultural statistics")]
    xs = [0.69, 2.89, 5.09, 7.29]
    for (h, sub), x in zip(boxes, xs):
        card(s, x, y, 2.0, 1.0, fill=TINT if x < 5 else RGBColor(0xF3, 0xF3, 0xF3))
        text(s, x + 0.08, y + 0.12, 1.84, 0.4, [h], size=14, bold=True, color=TEAL if x < 5 else INK,
             align=PP_ALIGN.CENTER)
        text(s, x + 0.08, y + 0.52, 1.84, 0.4, [sub], size=11, color=MUTED, align=PP_ALIGN.CENTER)
    text(s, 2.69, y + 0.28, 0.2, 0.4, ["×"], size=20, bold=True, color=MUTED, align=PP_ALIGN.CENTER)
    arrow(s, 4.9, y + 0.5, 5.08, y + 0.5)
    arrow(s, 7.1, y + 0.5, 7.28, y + 0.5)
    text(s, 0.69, y + 1.05, 4.4, 0.35, ["Our RO tests these two inputs where ZOA works"], size=12, bold=True,
         color=RED, align=PP_ALIGN.CENTER)
    # what the paper itself says
    text(s, 0.69, 3.1, 4.2, 1.9, [
        ("The authors' own warning", {"bold": True, "size": 14}),
        ("“Remote sensing data by itself lacks the level of accuracy necessary in a comprehensive food "
         "security assessment.”", {"italic": True, "size": 13}),
        ("Changing one input by ±20% changed the estimated losses by 12%.", {"size": 13}),
    ], size=13)
    text(s, 5.09, 3.1, 4.2, 1.9, [
        ("Same overlay, used today", {"bold": True, "size": 14}),
        ("WFP ADAM reports flooded cropland per county (flood map × NASA GFSAD30) with no accuracy "
         "statement.", {"size": 13}),
        ("Kerner et al. (2024) checked 11 cropland maps in 8 African countries, not South Sudan.",
         {"size": 13}),
    ], size=13)
    source(s, "Pacetti, Caporali & Rulli (2017), Advances in Water Resources 110:494–504 (Bangladesh 2007, "
              "Pakistan 2010). Kerner et al. (2024), Scientific Data, doi:10.1038/s41597-024-03306-z.")
    s.notes_slide.notes_text_frame.text = (
        "Pacetti et al. 2017 overlay a MODIS flood map (UNOSAT) on land-use maps, then convert lost crops into "
        "calories and water footprint. Quote from the discussion (p. 501). Sensitivity: water-depth threshold "
        "±20% -> 12% change in energy-content losses (abstract). ADAM: report of 2 May 2024, no accuracy "
        "statement. Kerner et al. 2024: maps agree unanimously on <0.5% of pixels; no single map optimal.")


def slide_neglected(s):
    items = [
        ("No ground truth", "Accuracy studies skip South Sudan (Kerner et al. 2024: 8 countries, none of them "
                            "South Sudan). We found no accuracy figure for flooded cropland there."),
        ("Clouds in the flood season", "Optical flood maps are blind in July–September, exactly when fields "
                                       "flood. Radar sees through clouds but is not used in the standard overlay."),
        ("Small, shifting fields", "Smallholder plots are small, move between years and lie fallow, so 10–30 m "
                                   "global maps easily miss them."),
        ("Numbers without error bars", "Tools publish hectares without uncertainty; CFSAM's harvested area is "
                                       "computed from population and household statistics, not measured."),
    ]
    pos = [(0.69, 1.55), (5.09, 1.55), (0.69, 3.35), (5.09, 3.35)]
    for i, ((h, b), (x, y)) in enumerate(zip(items, pos), 1):
        card(s, x, y, 4.22, 1.6)
        circle_num(s, x + 0.15, y + 0.17, 0.42, str(i))
        text(s, x + 0.7, y + 0.15, 3.4, 0.45, [h], size=15, bold=True, color=TEAL, anchor=MSO_ANCHOR.MIDDLE)
        text(s, x + 0.7, y + 0.6, 3.4, 0.95, [b], size=12)
    s.notes_slide.notes_text_frame.text = (
        "Sources: Kerner et al. 2024 Table 4 (countries). Optical blindness Jul-Sep: team's mask analysis and "
        "Downs 2023 (MODIS misses most flooding in South Sudan). CFSAM 2022 method: harvested area computed from "
        "population, household size, farming share and area per household (CFSAM 2022 summary, p. 2).")


def slide_costs(s):
    stats = [("35×", "less flooded cropland in the standard overlay than CFSAM reports",
              "South Sudan, Jun–Nov 2022: flood masks × ASAP ≈ 3,700 ha vs 130,000 ha damaged"),
             ("0 ha", "flooded cropland ever detected in Bor South",
              "flood masks × ASAP, every season checked, in an area ZOA names as flood-prone (Q&A 2, slide 3)"),
             ("70 / 77", "counties where the ASAP cropland map is below CFSAM's harvested area",
              "national: ≈407,000 ha mapped vs 1.03–1.18 million ha harvested")]
    for i, (big, what, detail) in enumerate(stats):
        x = 0.69 + i * 2.93
        text(s, x, 1.5, 2.75, 0.8, [big], size=40, bold=True, color=RED)
        text(s, x, 2.3, 2.75, 0.8, [what], size=13, bold=True)
        text(s, x, 3.05, 2.75, 0.75, [detail], size=10, color=MUTED)
    card(s, 0.69, 3.85, 8.62, 1.25)
    text(s, 0.85, 3.92, 8.3, 0.35, ["What that means for anticipatory action"], size=14, bold=True, color=TEAL)
    bullets(s, 0.85, 4.28, 8.3, 0.85, [
        "A trigger on “flood inundates cropland” would rarely fire, even in seasons when harvests are lost.",
        "Aid follows the data, not the damage; and figures without uncertainty look more precise than they are.",
    ], size=12, space=3)
    s.notes_slide.notes_text_frame.text = (
        "Numbers from archive/external_reviews/superior_ro_search_report.md: check 2 (masks x ASAP, 2022 ~3,700 ha "
        "nationally; 0-470 ha per ZOA county-season; 0 ha in Bor South every year) and check 1 (ASAP vs CFSAM, "
        "70 of 77 matched counties). CFSAM 2022 summary p. 1: 130,000 ha of cultivated land damaged. CFSAM is "
        "itself a modelled estimate, so the gap shows the size of the uncertainty, not the true loss.")


def slide_progress(s):
    steps = [("Collect", "11 public datasets: 7 cropland maps, flood, terrain, buildings"),
             ("Align", "one Sentinel-2-aligned 10 / 30 m grid for Aweil and Bor South"),
             ("Check", "codes verified on the pixels; manifest with sources and checksums"),
             ("Compare", "map-to-map agreement per county"),
             ("Stratify", "strata for the blind sample (maps saying crop × distance to homes)")]
    for i, (h, b) in enumerate(steps):
        x = 0.69 + i * 1.75
        circle_num(s, x + 0.55, 1.65, 0.5, str(i + 1))
        if i < 4:
            arrow(s, x + 1.1, 1.9, x + 2.25, 1.9)
        text(s, x, 2.25, 1.62, 0.4, [h], size=15, bold=True, color=TEAL, align=PP_ALIGN.CENTER)
        text(s, x, 2.65, 1.62, 1.3, [b], size=11, align=PP_ALIGN.CENTER)
    card(s, 0.69, 4.0, 8.62, 1.05)
    text(s, 0.85, 4.07, 8.3, 0.95, [
        ("Done this sprint: steps 1–4 for both areas, scripted and re-runnable.", {"bold": True, "size": 13}),
        ("Next: Sentinel-1 radar flood layer for 2022, fix the strata, draw the sample, 100-point labelling pilot.",
         {"size": 13}),
    ], size=13)
    s.notes_slide.notes_text_frame.text = (
        "Code: processing_data/cropland/ (download_*.py, derived_cropland_layers.py, cropland_agreement.py, "
        "check_cropland_rasters.py). Documentation: data/cropland/README.md, MANIFEST.csv, CHECKS.md.")


def slide_demo(s, fig_path):
    s.shapes.add_picture(str(fig_path), Inches(0.6), Inches(1.5), height=Inches(3.3))   # 2:1 figure -> 6.6" wide
    text(s, 0.69, 5.07, 8.62, 0.35, [
        "Same 6 × 6 km near Turalei (Twic county), outside our study counties so the labelling stays blind. "
        "Gold = map says crop."], size=10, color=MUTED)
    text(s, 7.5, 1.9, 1.8, 0.6, ["0.2%"], size=30, bold=True, color=RED)
    text(s, 7.5, 2.5, 1.8, 0.4, ["to"], size=13, color=MUTED)
    text(s, 7.5, 2.85, 1.8, 0.6, ["48.9%"], size=30, bold=True, color=RED)
    text(s, 7.5, 3.5, 1.8, 1.3, ["of the same window called crop, depending on which map you ask"], size=12)
    s.notes_slide.notes_text_frame.text = (
        "Seven maps, seven answers for one place: from 0.2% (WorldCover, Esri) to 48.9% (GFSAD) of the window. "
        "Deliberately outside Aweil and Bor South: labellers must not see the maps in the study area. "
        "Figure: data/cropland/demo/, script processing_data/cropland/demo_map_panels.py.")


def slide_chart(s):
    df = aweil_shares()
    cd = CategoryChartData()
    cd.categories = list(df.source)
    cd.add_series("Share of the 5 Aweil counties mapped as crop (%)", [round(v, 2) for v in df.share])
    gf = s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(0.55), Inches(1.45), Inches(5.6), Inches(3.7), cd)
    ch = gf.chart
    ch.has_legend = False
    ch.has_title = True
    ch.chart_title.text_frame.text = "Cropland in the 5 Aweil counties (% of land)"
    ch.chart_title.text_frame.paragraphs[0].runs[0].font.size = Pt(12)
    ch.chart_title.text_frame.paragraphs[0].runs[0].font.bold = True
    plot = ch.plots[0]
    plot.gap_width = 60
    plot.has_data_labels = True
    dl = plot.data_labels
    dl.number_format = '0.0"%"'
    dl.number_format_is_linked = False
    dl.position = XL_LABEL_POSITION.OUTSIDE_END
    dl.font.size = Pt(10)
    ser = plot.series[0]
    for i, name in enumerate(df.source):
        pt = ser.points[i]
        pt.format.fill.solid()
        pt.format.fill.fore_color.rgb = GREY if name.startswith("CFSAM") else (TEAL if name.startswith("ASAP") else GOLD)
    va = ch.value_axis
    va.has_major_gridlines = False
    va.visible = False
    va.maximum_scale = 6.5
    ca = ch.category_axis
    ca.tick_labels.font.size = Pt(10)
    ca.format.line.color.rgb = RGBColor(0xBF, 0xBF, 0xBF)
    # callouts
    calls = [("1 ha", "where all four main maps (2021–22) agree on crop, in 3 million ha"),
             ("≤ 25%", "best overlap of the crop areas of any two maps"),
             ("5×", "CFSAM's modelled harvested area vs the typical map")]
    for i, (big, what) in enumerate(calls):
        y = 1.55 + i * 1.2
        text(s, 6.35, y, 2.95, 0.55, [big], size=26, bold=True, color=RED)
        text(s, 6.35, y + 0.5, 2.95, 0.65, [what], size=11)
    source(s, "Crop share from 10 m pixels, counties' own area. CFSAM = 2022 main-harvest cereal area, a modelled "
              "estimate (context, not a reference). Agreement is between maps, not accuracy.")
    s.notes_slide.notes_text_frame.text = (
        "Data: data/cropland/agreement/ (county_crop_share.csv, pairwise_agreement.csv, strata_n_maps.csv). "
        "Four main maps = WorldCereal 2021, WorldCover 2021, Esri 2022, Dynamic World Jun-Nov 2022. Best "
        "pairwise crop overlap (Jaccard) 25% (WorldCereal vs WorldCover, one county); kappa at most 0.39. "
        "Typical map ~1% vs CFSAM 5.4%. Consequence for the design: 'all maps agree' cannot be a stratum; "
        "the sample must also cover the 'no map says crop' land, where most real cropland may be.")


def slide_lessons(s):
    rows = [("Took crops = 4 in Esri from the catalogue example", "On the pixels crops = 5; 4 is flooded vegetation"),
            ("Trusted documented no-data (DE Africa: 0)", "The files use 255; 0 means “not crop”"),
            ("GFSAD code order differs between NASA pages", "The user guide decides: 2 = cropland"),
            ("Judged maps in zoomed-out previews", "Previews average class codes; use full resolution only"),
            ("Expected one map to work everywhere", "Dynamic World calls 23% of Bor South crop (it is grassland)")]
    tbl = s.shapes.add_table(len(rows) + 1, 2, Inches(0.69), Inches(1.5), Inches(8.62), Inches(3.1)).table
    tbl.columns[0].width = Inches(4.2)
    tbl.columns[1].width = Inches(4.42)
    for j, h in enumerate(["Mistake or trap", "What we learned"]):
        c = tbl.cell(0, j)
        c.text = h
        c.fill.solid()
        c.fill.fore_color.rgb = TEAL
        r = c.text_frame.paragraphs[0].runs[0]
        r.font.bold = True
        r.font.size = Pt(13)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    for i, (a, b) in enumerate(rows, 1):
        for j, v in enumerate((a, b)):
            c = tbl.cell(i, j)
            c.text = v
            c.fill.solid()
            c.fill.fore_color.rgb = TINT if i % 2 else RGBColor(0xFF, 0xFF, 0xFF)
            r = c.text_frame.paragraphs[0].runs[0]
            r.font.size = Pt(12)
            r.font.color.rgb = INK
    text(s, 0.69, 4.72, 8.62, 0.45, [("Lesson: check every class code on the data itself, and write it down "
                                      "(manifest + README).", {"bold": True, "color": RED})], size=13)
    s.notes_slide.notes_text_frame.text = (
        "All five were caught by the checks in processing_data/cropland/check_cropland_rasters.py and are "
        "documented in data/cropland/README.md. The Esri trap alone would have mapped flooded vegetation as crop.")


def slide_stakeholder(s):
    quotes = [("Harvest failure is the priority impact", "slide 6", "and ≈80% of rural communities live "
                                                                      "mainly from agriculture (slide 33)"),
              ("“The dataset might not capture agricultural land”", "slide 33",
               "ZOA already suspects the undercount our numbers show"),
              ("Triggers like a flood “inundating cropland”", "slides 31, 40",
               "only usable if the cropland and flood data are reliable"),
              ("Show the reliability, not a perfect forecast", "slides 9–11, 32",
               "a rigorous small sub-problem beats a superficial full one (slides 5, 29)")]
    for i, (q, ref, why) in enumerate(quotes):
        x = 0.69 + (i % 2) * 4.4
        y = 1.55 + (i // 2) * 1.8
        card(s, x, y, 4.22, 1.6)
        text(s, x + 0.18, y + 0.12, 3.9, 0.7, [q], size=14, bold=True, color=TEAL)
        text(s, x + 0.18, y + 0.82, 3.9, 0.3, [ref], size=10, color=MUTED)
        text(s, x + 0.18, y + 1.08, 3.9, 0.5, [why], size=11)
    source(s, "Stakeholder Q&A 2 (ZOA / Zero Hunger Lab, 21 Sep 2026). Our RO covers ZOA's two areas: the Lol "
              "river in Aweil and Bor South (slide 3).")
    s.notes_slide.notes_text_frame.text = (
        "Progress with the stakeholders: the Q&A 2 answers fixed the areas, the scale (community) and what ZOA "
        "wants from data (reliability and risks). The cropland audit answers the undercount concern directly. "
        "Quotes as cited in the team's RO search report; check wording against the slides before presenting.")


# --- SLE slides ----------------------------------------------------------------------
# Lecture page references follow course_content/SLE_aspects_capstone_extracted_new.md
# (L1 = Petročnik, L2/L3 = van Maanen). They are leads for the essay, not essay text.

SLE = "Societal-Legal-Ethical Angle (Essay): "


def sle_figure() -> Path:
    """Satellite image next to the map with the least and the most cropland (same window as the demo)."""
    rgb = rasterio.open(DEMO / "sentinel2_oct_dec_2022_rgb.tif").read().astype(float)
    rgb = np.clip(np.moveaxis(rgb, 0, -1) / 3000.0, 0, 1)
    fig, axes = plt.subplots(1, 3, figsize=(12.9, 4.6), constrained_layout=True)
    axes[0].imshow(rgb)
    axes[0].set_title("Satellite image (late 2022)", fontsize=17)
    cmap = ListedColormap(["#ece7dc", "#d4a017", "#ffffff"])
    for ax, (name, f, code) in zip(axes[1:], [("WorldCover 2021", "worldcover.tif", 40),
                                              ("GFSAD 2015", "gfsad2015.tif", 2)]):
        a = rasterio.open(DEMO / f).read(1)
        share = 100 * (a[a != 255] == code).mean()
        ax.imshow(np.where(a == 255, 2, (a == code).astype(int)), cmap=cmap, vmin=0, vmax=2,
                  interpolation="nearest")
        ax.set_title(f"{name}: {share:.1f}% crop", fontsize=17)
    for ax in axes:
        ax.set_xticks([])
        ax.set_yticks([])
    path = FIG / "slide_sle_two_maps_turalei.png"
    fig.savefig(path, dpi=150)
    plt.close(fig)
    return path


def aweil_rule_areas() -> dict[str, float]:
    """Aweil cropland (ha) under three agreement rules, from the n-maps-say-crop strata."""
    s = pd.read_csv(ROOT / "data" / "cropland" / "agreement" / "strata_n_maps.csv")
    s = s[(s.aoi == "aweil") & (s.n_maps_crop != "no data in >= 1 map")].copy()
    s["n"] = s.n_maps_crop.astype(int)
    four, seven = s[s.map_set == "4 primary maps"], s[s.map_set == "all 7 maps"]
    return {"all4": four[four.n == 4].area_ha.sum(), "any4": four[four.n >= 1].area_ha.sum(),
            "any7": seven[seven.n >= 1].area_ha.sum()}


def slide_sle_problem(s):
    text(s, 0.69, 1.5, 4.4, 0.35, ["The problem"], size=15, bold=True, color=TEAL)
    bullets(s, 0.69, 1.85, 4.4, 1.6, [
        "ZOA's example trigger: a flood “inundating cropland” (Q&A 2, slides 31, 40).",
        "That number comes from a public flood map laid over a public cropland map.",
        "Nobody has checked those maps along the Lol river (Aweil) or in Bor South.",
    ], size=12, space=4)
    text(s, 0.69, 3.4, 4.4, 0.35, ["Why it matters"], size=15, bold=True, color=TEAL)
    bullets(s, 0.69, 3.75, 4.4, 1.4, [
        "≈80% of rural communities live mainly from farming, mostly smallholders (slide 33).",
        "Harvest failure is ZOA's priority flood impact (slide 6).",
        "ZOA itself warns: “the dataset might not capture agricultural land” (slide 33).",
    ], size=12, space=4)
    card(s, 5.3, 1.5, 4.01, 3.6)
    text(s, 5.48, 1.62, 3.7, 0.35, ["SLE question (draft, for the team)"], size=14, bold=True, color=TEAL)
    text(s, 5.48, 2.0, 3.7, 1.3, ["Whose farmland counts? How the choice of a satellite map decides which "
                                  "flood-hit farmers anticipatory action can see."], size=15, bold=True)
    text(s, 5.48, 3.3, 3.7, 1.75, [
        ("Lenses from the lectures", {"bold": True, "size": 11, "color": MUTED}),
        ("“Technical choices = SLE choices” (L1 p. 5)", {"size": 11}),
        ("“Whose way of ‘seeing’ counts?” (Jasanoff, L1 p. 7)", {"size": 11}),
        ("“Which communities? Who? Where?” (Sen, L2 p. 11)", {"size": 11}),
    ], size=11)
    source(s, "Stakeholder Q&A 2 (ZOA / Zero Hunger Lab, 21 Sep 2026). L1 = SLE lecture 1 (Petročnik), "
              "L2 / L3 = SLE lectures 2 and 3 (van Maanen).")
    s.notes_slide.notes_text_frame.text = (
        "Same problem as the methodological slides, but asked as an SLE question: the choice of map is not "
        "neutral, it decides who is visible. The question is a draft for the team; the essay text must be "
        "the team's own. Check lecture page numbers against the lecture PDFs before citing them.")


def slide_sle_seen(s, fig_path, areas):
    s.shapes.add_picture(str(fig_path), Inches(0.69), Inches(1.42), width=Inches(8.62))  # 2.8:1 -> 3.07" high
    text(s, 0.69, 4.5, 4.2, 0.7, [
        ("Same fields, near Turalei (outside our study area): one map sees almost no farms, the other sees "
         "half the land.", {"size": 11}),
    ], size=11)
    text(s, 5.09, 4.5, 4.22, 0.7, [
        (f"Aweil counties: {areas['all4']:,.0f} ha of cropland if all four main maps must agree, "
         f"≈{round(areas['any4'], -3):,.0f} ha if one is enough, ≈{round(areas['any7'], -3):,.0f} ha if any of "
         f"seven counts.",
         {"size": 11, "bold": True, "color": RED}),
    ], size=11)
    source(s, "Gold = map says crop. Reading: data is socially constructed, not simply “seen” (L1 p. 7); "
              "unrepresentative data create blind spots (L1 p. 10). Agreement between maps is not accuracy.")
    s.notes_slide.notes_text_frame.text = (
        "The crop-mask result from the method part, told as an SLE point: which farmer 'exists' depends on the "
        "map picked, and that pick is usually made silently by the tool. Window: 6 x 6 km near Turalei (Twic), "
        "kept outside Aweil and Bor South so the labelling stays blind. Aweil numbers from "
        "data/cropland/agreement/strata_n_maps.csv (four main maps = WorldCereal 2021, WorldCover 2021, "
        "Esri 2022, Dynamic World Jun-Nov 2022; seven = plus GLAD 2019, DE Africa 2019, GFSAD 2015).")


def slide_sle_neglected(s):
    card(s, 0.69, 1.5, 4.22, 2.75)
    text(s, 0.87, 1.6, 3.9, 0.35, ["The paper we build on"], size=14, bold=True, color=TEAL)
    text(s, 0.87, 1.98, 3.9, 2.25, [
        ("Pacetti, Caporali & Rulli (2017): flood map × land-use map → lost crops and calories "
         "(Bangladesh 2007, Pakistan 2010).", {"size": 12}),
        ("Their own warning: “Remote sensing data by itself lacks the level of accuracy necessary in a "
         "comprehensive food security assessment.”", {"size": 12, "italic": True}),
        ("WFP ADAM uses the same overlay for South Sudan, with no accuracy statement.", {"size": 12}),
    ], size=12)
    text(s, 5.09, 1.5, 4.22, 0.35, ["Why the problem is neglected"], size=14, bold=True, color=TEAL)
    bullets(s, 5.09, 1.88, 4.22, 2.4, [
        "No ground truth: accuracy studies skip South Sudan (Kerner et al. 2024: 8 countries, not this one).",
        "Clouds hide the fields in July–September, exactly when they flood.",
        "Smallholder plots are small, move and lie fallow: global maps miss them.",
        "Hectares are published without error bars, so the gap stays invisible.",
    ], size=12, space=4)
    card(s, 0.69, 4.4, 8.62, 0.72, fill=RGBColor(0xF3, 0xF3, 0xF3))
    text(s, 0.85, 4.45, 8.3, 0.65, [
        ("Leads for the essay: Who writes what, and for what reason? (L3 p. 18). Do exact-looking hectares serve "
         "reporting to donors more than decisions in the field? (logic of audit, L2 p. 4)", {"size": 12}),
    ], size=12, anchor=MSO_ANCHOR.MIDDLE)
    source(s, "Pacetti et al. (2017), Advances in Water Resources 110:494–504, quote p. 501. "
              "Kerner et al. (2024), Scientific Data, doi:10.1038/s41597-024-03306-z.")
    s.notes_slide.notes_text_frame.text = (
        "Neglect has technical reasons (clouds, small fields, no ground truth) and possibly institutional ones: "
        "the lectures ask who produces knowledge and for whom. The audit point is a question to explore, not a "
        "finding. ADAM: report of 2 May 2024, no accuracy statement.")


def slide_sle_costs(s):
    items = [
        ("Farmers left out", "A field missing from the map has no flooding and no household to target. Bor South: "
                             "0 ha flooded cropland ever detected, yet its roads close 77–92% of weeks in Sep–Dec.",
         "Targeting (Chavez-Gonzalez et al., L3 p. 16); blind spots (L1 p. 10)"),
        ("Triggers that stay silent", "2022: the standard overlay finds ≈3,700 ha of flooded cropland in South "
                                      "Sudan; CFSAM reports 130,000 ha damaged. The trigger sees ≈3%.",
         "Timing and activity: acting too late or not at all (L3 p. 16)"),
        ("False certainty", "Hectares without error bars look exact. ZOA asked for the risks of the data so they "
                            "can choose between misses and false alarms (Q&A 2, slides 2, 9–11).",
         "Data is not self-explanatory (Letz & Maxwell, L3 p. 17)"),
        ("Hectares are not hunger", "Flooded cropland is one factor. The same flood hits IDPs, hosts and "
                                    "pastoralists differently (Q&A 2, slide 38).",
         "Do not use one factor; entitlements (Sen, L2 pp. 8–10)"),
    ]
    pos = [(0.69, 1.45), (5.09, 1.45), (0.69, 3.3), (5.09, 3.3)]
    for i, ((h, b, lec), (x, y)) in enumerate(zip(items, pos), 1):
        card(s, x, y, 4.22, 1.75)
        circle_num(s, x + 0.15, y + 0.14, 0.4, str(i), fill=RED)
        text(s, x + 0.68, y + 0.12, 3.45, 0.42, [h], size=14, bold=True, color=TEAL, anchor=MSO_ANCHOR.MIDDLE)
        text(s, x + 0.18, y + 0.58, 3.95, 0.85, [b], size=11)
        text(s, x + 0.18, y + 1.43, 3.95, 0.3, [lec], size=9, italic=True, color=MUTED)
    s.notes_slide.notes_text_frame.text = (
        "Who bears the error: a missed field falls on the smallholder, a false alarm on the budget. Numbers: "
        "flood masks x ASAP (2022, national, ~3,700 ha) vs CFSAM 2022 summary p. 1 (130,000 ha of cultivated land "
        "damaged); CFSAM is itself modelled, so ~3% shows the size of the gap, not the true loss. Bor South roads: "
        "Logistics Cluster maps 2022-2026, archive/roads_flood/ROADS_FLOOD_EDA.md.")


def slide_sle_lessons(s):
    rows = [("We started from what a model can predict (a flood forecast)",
             "ZOA decides about crops and people; tools can change the problem (L3 p. 11). We check the data first."),
            ("We took one cropland map as the truth",
             "Seven maps give seven answers: a map is a way of seeing (L1 p. 7)"),
            ("We used CFSAM harvested area as the reference",
             "It is modelled from population data: “one source = no source” (L3 p. 22)"),
            ("We took “flood” to mean what the flood mask detects",
             "Bor South roads close while masks see nothing: “what is a flood?” (L2 p. 10)")]
    tbl = s.shapes.add_table(len(rows) + 1, 2, Inches(0.69), Inches(1.5), Inches(8.62), Inches(3.0)).table
    tbl.columns[0].width = Inches(3.9)
    tbl.columns[1].width = Inches(4.72)
    for j, h in enumerate(["Mistake", "What we learned"]):
        c = tbl.cell(0, j)
        c.text = h
        c.fill.solid()
        c.fill.fore_color.rgb = TEAL
        r = c.text_frame.paragraphs[0].runs[0]
        r.font.bold = True
        r.font.size = Pt(13)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    for i, (a, b) in enumerate(rows, 1):
        for j, v in enumerate((a, b)):
            c = tbl.cell(i, j)
            c.text = v
            c.fill.solid()
            c.fill.fore_color.rgb = TINT if i % 2 else RGBColor(0xFF, 0xFF, 0xFF)
            r = c.text_frame.paragraphs[0].runs[0]
            r.font.size = Pt(12)
            r.font.color.rgb = INK
    text(s, 0.69, 4.62, 8.62, 0.5, [("Lesson: report every number with its uncertainty and its blind spots, and "
                                     "write down our assumptions (Chavez-Gonzalez et al., L3 p. 16).",
                                     {"bold": True, "color": RED})], size=13)
    s.notes_slide.notes_text_frame.text = (
        "Team: check that these four match your own experience before presenting. They come from the RO "
        "history: the forecast was dropped from the RO (27 Sep plan), CFSAM turned out to be modelled, the "
        "cropland maps disagree, and the road data showed closures in Bor South the flood masks miss.")


def main() -> None:
    FIG.mkdir(parents=True, exist_ok=True)
    prs = Presentation(str(TEMPLATE))
    slides = list(prs.slides)
    tpl_title = copy.deepcopy(slides[3].shapes.title._element.txBody)
    layout = next(l for l in prs.slide_layouts if l.name == "Title and Content")

    # existing slides 4, 5, 6, 8, 9: replace the placeholder bodies (8 and 9 also get the two-tone title)
    s4, s5, s6, s8, s9 = slides[3], slides[4], slides[5], slides[7], slides[8]
    for s in (s4, s5, s6, s8, s9):
        drop_body(s)
    slide_progress(s4)
    slide_lessons(s5)
    slide_stakeholder(s6)
    set_title(s8, tpl_title, SLE, "Progress: the problem as an SLE question")
    slide_sle_problem(s8)
    set_title(s9, tpl_title, SLE, "Mistakes and what we have learned")
    slide_sle_lessons(s9)

    new = {}
    for key, l2, fn in [("problem", "Problem definition & research objective", slide_problem),
                        ("paper", "The method we build on", slide_paper),
                        ("neglected", "Why the problem is neglected", slide_neglected),
                        ("costs", "What unreliable data costs", slide_costs),
                        ("demo", "First results: seven maps, seven answers", slide_demo),
                        ("chart", "First results: cropland per map in Aweil", slide_chart)]:
        s = prs.slides.add_slide(layout)
        drop_body(s)
        set_title(s, tpl_title, "Methodological Angle (Poster/Report): ", l2)
        if key == "demo":
            fn(s, demo_figure())
        else:
            fn(s)
        new[key] = s

    areas = aweil_rule_areas()
    for key, l2, fn in [("sle_seen", "Result: the map decides who is seen",
                         lambda s: slide_sle_seen(s, sle_figure(), areas)),
                        ("sle_neglected", "Prior paper and why it is neglected",
                         slide_sle_neglected),
                        ("sle_costs", "What unreliable data costs people", slide_sle_costs)]:
        s = prs.slides.add_slide(layout)
        drop_body(s)
        set_title(s, tpl_title, SLE, l2)
        fn(s)
        new[key] = s

    # order: 1-3, problem, paper, neglected, costs, 4 (progress), demo, chart, 5 (lessons), 6-7 (stakeholder),
    # 8 (SLE problem), SLE seen / neglected / costs, 9 (SLE lessons), 10
    sld = prs.slides._sldIdLst
    ids = list(sld)
    by_id = {int(e.get("id")): e for e in ids}           # <p:sldId id=...> == slide.slide_id
    order = (slides[:3] + [new["problem"], new["paper"], new["neglected"], new["costs"], s4, new["demo"],
                           new["chart"], s5] + slides[5:7]
             + [s8, new["sle_seen"], new["sle_neglected"], new["sle_costs"], s9] + slides[9:])
    assert len(order) == len(ids)
    wanted = [s.slide_id for s in order]                 # read before the list is emptied
    for e in ids:
        sld.remove(e)
    for sid in wanted:
        sld.append(by_id[sid])
    prs.save(str(OUT))
    print(f"wrote {OUT} ({len(order)} slides)")


if __name__ == "__main__":
    main()
