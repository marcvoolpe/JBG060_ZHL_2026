"""Build the deck for the methodological feedback session (1 Oct 2026): story, RO, method, questions.

Reads course_content/Template_sprint_review.pptx (never modified) and writes
deliverables/method_feedback/Method_feedback_cropland.pptx. Style helpers, the Turalei figure and the
Aweil crop shares come from deliverables/sprint_review/build_sprint_review_deck.py, so both decks match.
Method content follows deliverables/CROPLAND_PLAN.html (plan of 28 Sep 2026).

Usage (repo root, global Python 3.13):
    python deliverables/method_feedback/build_method_feedback_deck.py
"""

from __future__ import annotations

import copy
import importlib.util
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[2]
TEMPLATE = ROOT / "course_content" / "Template_sprint_review.pptx"
OUT = ROOT / "deliverables" / "method_feedback" / "Method_feedback_cropland.pptx"
TURALEI = ROOT / "deliverables" / "sprint_review" / "slide_demo_maps_turalei.png"

_spec = importlib.util.spec_from_file_location(
    "sprint_deck", ROOT / "deliverables" / "sprint_review" / "build_sprint_review_deck.py")
sd = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(sd)
text, bullets, card, circle_num, arrow, source = sd.text, sd.bullets, sd.card, sd.circle_num, sd.arrow, sd.source
RED, INK, MUTED, TINT, TEAL, GOLD, GREY = sd.RED, sd.INK, sd.MUTED, sd.TINT, sd.TEAL, sd.GOLD, sd.GREY
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT = RGBColor(0xF3, 0xF3, 0xF3)


def table(slide, x, y, w, col_w, rows, header, size=11, head_size=12, red_rows=()):
    """Native table: teal header, alternating tint rows; rows listed in red_rows get red bold text."""
    shape = slide.shapes.add_table(len(rows) + 1, len(header), Inches(x), Inches(y), Inches(w),
                                   Inches(0.3 * (len(rows) + 1)))
    tbl = shape.table
    for j, cw in enumerate(col_w):
        tbl.columns[j].width = Inches(cw)
    for i, row in enumerate([header] + rows):
        for j, v in enumerate(row):
            c = tbl.cell(i, j)
            c.text = v
            c.margin_left = c.margin_right = Inches(0.06)
            c.margin_top = c.margin_bottom = Inches(0.03)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            c.fill.solid()
            c.fill.fore_color.rgb = TEAL if i == 0 else (TINT if i % 2 else WHITE)
            p = c.text_frame.paragraphs[0]
            if j > 0 and all(ch in "0123456789.±– " for ch in v):
                p.alignment = PP_ALIGN.CENTER
            r = p.runs[0]
            r.font.size = Pt(head_size if i == 0 else size)
            r.font.bold = i == 0 or (i - 1) in red_rows
            r.font.color.rgb = WHITE if i == 0 else (RED if (i - 1) in red_rows else INK)
    return tbl


def notes(slide, s):
    slide.notes_slide.notes_text_frame.text = s


# --- slides ------------------------------------------------------------------------

def slide_title(s):
    boxes = [sh for sh in s.shapes if sh.has_text_frame]
    runs = boxes[0].text_frame.paragraphs[0].runs
    runs[0].text = "Methodological Feedback"
    runs[1].text = "Counting South Sudan's cropland"
    runs[1].font.size = Pt(34)
    notes(s, "Goal of the session: confirm the research objective, the method and the story. "
             "We end with four questions for feedback.")


def slide_why(s):
    y = 1.55
    boxes = [("Flood map", "NASA MODIS masks"), ("Cropland map", "ASAP, WorldCover, …"),
             ("Flooded cropland", "hectares per county"), ("AA trigger", "flood “inundating cropland”")]
    xs = [0.69, 2.89, 5.09, 7.29]
    for i, ((h, sub), x) in enumerate(zip(boxes, xs)):
        card(s, x, y, 2.0, 1.0, fill=TINT if i < 3 else LIGHT)
        text(s, x + 0.08, y + 0.12, 1.84, 0.4, [h], size=14, bold=True, color=TEAL if i < 3 else INK,
             align=PP_ALIGN.CENTER)
        text(s, x + 0.08, y + 0.52, 1.84, 0.4, [sub], size=11, color=MUTED, align=PP_ALIGN.CENTER)
    text(s, 2.69, y + 0.28, 0.2, 0.4, ["×"], size=20, bold=True, color=MUTED, align=PP_ALIGN.CENTER)
    arrow(s, 4.9, y + 0.5, 5.08, y + 0.5)
    arrow(s, 7.1, y + 0.5, 7.28, y + 0.5)
    text(s, 2.89, y + 1.05, 2.0, 0.35, ["the weak link"], size=12, bold=True, color=RED, align=PP_ALIGN.CENTER)

    stats = [("35×", "less flooded cropland than the FAO/WFP estimate",
              "South Sudan 2022: flood masks × ASAP ≈ 3,700 ha; CFSAM reports 130,000 ha damaged"),
             ("0 ha", "flooded cropland ever detected in Bor South",
              "every season checked, in an area ZOA names as flood-prone"),
             ("“…might not capture agricultural land”", "ZOA about the course cropland data",
              "Stakeholder Q&A 2, slide 33")]
    for i, (big, what, detail) in enumerate(stats):
        x = 0.69 + i * 2.93
        big_size = 36 if len(big) < 6 else 16
        text(s, x, 3.05, 2.75, 0.75, [big], size=big_size, bold=True, color=RED,
             anchor=MSO_ANCHOR.BOTTOM, italic=len(big) > 6)
        text(s, x, 3.85, 2.75, 0.55, [what], size=13, bold=True)
        text(s, x, 4.45, 2.75, 0.65, [detail], size=10, color=MUTED)
    notes(s, "ZOA's example trigger is a flood 'inundating cropland' (Q&A 2, slides 31 and 40). The number behind "
             "it comes from laying a flood map over a cropland map (Pacetti et al. 2017; WFP ADAM does the same). "
             "If the cropland map misses the fields, the trigger stays silent. Numbers: flood masks x ASAP, "
             "Jun-Nov 2022, national (~3,700 ha) vs CFSAM 2022 (130,000 ha of cultivated land damaged). CFSAM is "
             "itself modelled, so 35x shows the size of the gap, not the true loss. Bor South: 0 ha in every season.")


def slide_seven_maps(s):
    s.shapes.add_picture(str(TURALEI), Inches(0.6), Inches(1.5), height=Inches(3.3))   # 2:1 -> 6.6" wide
    text(s, 7.45, 1.45, 1.9, 0.5, ["0.2%"], size=28, bold=True, color=RED)
    text(s, 7.45, 1.93, 1.9, 0.3, ["to"], size=11, color=MUTED)
    text(s, 7.45, 2.18, 1.9, 0.5, ["48.9%"], size=28, bold=True, color=RED)
    text(s, 7.45, 2.7, 1.9, 0.5, ["crop in the same window"], size=11)
    text(s, 7.45, 3.3, 1.9, 1.5, [
        ("By eye, only GFSAD covers the cleared land, and it is the oldest map (2015).",
         {"size": 11, "color": TEAL, "bold": True}),
        ("Not measured yet: that is what our labelled points are for.", {"size": 10, "color": MUTED})],
        size=11)
    text(s, 0.69, 5.05, 8.62, 0.35, [
        "Same 6 × 6 km window near Turalei (Twic county), outside our study counties so the labelling stays blind. "
        "Gold = map says crop."], size=10, color=MUTED)
    notes(s, "Seven maps, seven answers for one place: from 0.2% (WorldCover 2021, Esri 2022) to 48.9% (GFSAD "
             "2015). Looking at the Sentinel-2 image, GFSAD is the only map that covers the cleared land along the "
             "river, but that is our eye, not a measurement; GFSAD also paints whole blocks at 30 m and may include "
             "grass and fallow. The window is deliberately outside Aweil and Bor South so labellers never see the "
             "maps in the study area. Script: processing_data/cropland/demo_map_panels.py.")


def slide_disagree(s):
    df = sd.aweil_shares()
    cd = CategoryChartData()
    cd.categories = list(df.source)
    cd.add_series("Share of the 5 Aweil counties mapped as crop (%)", [round(v, 2) for v in df.share])
    gf = s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(0.55), Inches(1.45), Inches(5.6), Inches(3.7), cd)
    ch = gf.chart
    ch.has_legend = False
    ch.has_title = True
    ch.chart_title.text_frame.text = "Cropland in the 5 Aweil counties (% of land)"
    ch.chart_title.text_frame.paragraphs[0].runs[0].font.size = Pt(12)
    ch.chart_title.text_frame.paragraphs[0].runs[0].font.bold = True
    plot = ch.plots[0]
    plot.gap_width = 60
    plot.has_data_labels = True
    dl = plot.data_labels
    dl.number_format = '0.0"%"'
    dl.number_format_is_linked = False
    dl.position = XL_LABEL_POSITION.OUTSIDE_END
    dl.font.size = Pt(10)
    ser = plot.series[0]
    for i, name in enumerate(df.source):
        pt = ser.points[i]
        pt.format.fill.solid()
        pt.format.fill.fore_color.rgb = GREY if name.startswith("CFSAM") else (TEAL if name.startswith("ASAP") else GOLD)
    va = ch.value_axis
    va.has_major_gridlines = False
    va.visible = False
    va.maximum_scale = 6.5
    ca = ch.category_axis
    ca.tick_labels.font.size = Pt(10)
    ca.format.line.color.rgb = RGBColor(0xBF, 0xBF, 0xBF)

    calls = [("1 ha", "where the four newest maps (2021–22) all agree on crop, out of 3 million ha. "
                      "Best overlap of any two maps: ≤ 25%."),
             ("5×", "more crop in the FAO/WFP CFSAM 2022 report than in the typical map. CFSAM is modelled "
                    "from household statistics, not mapped."),
             ("2015–22", "map years. GFSAD, DE Africa and GLAD are 2019 or older; ASAP merges older maps. "
                         "The newest three find the least crop.")]
    for i, (big, what) in enumerate(calls):
        y = 1.5 + i * 1.22
        text(s, 6.35, y, 2.95, 0.5, [big], size=24, bold=True, color=RED)
        text(s, 6.35, y + 0.46, 2.95, 0.75, [what], size=10.5)
    source(s, "Crop share from 10 m pixels over the counties' area. CFSAM = 2022 main-harvest cereal area, a "
              "modelled estimate (context, not a reference). Agreement between maps is not accuracy.")
    notes(s, "Data: data/cropland/agreement/ (county_crop_share.csv, pairwise_agreement.csv, strata_n_maps.csv) and "
             "county_cropland_areas.csv. Four newest maps = WorldCereal 2021, WorldCover 2021, Esri 2022, Dynamic "
             "World Jun-Nov 2022. Best pairwise crop overlap (Jaccard) 25% (WorldCereal vs WorldCover, one county); "
             "kappa at most 0.39. If one map saying crop is enough, Aweil has ~47,000 ha; if any of seven counts, "
             "~140,000 ha. Map years: GLAD '2019' covers 2016-19; ASAP uses GlobCover in most of Africa (Kerner et "
             "al. 2024). The newest three (Dynamic World, WorldCover, Esri) all give 0.4% or less.")


def slide_kerner(s):
    steps = [("Random sample, not the map",
              "Random points per country. In Mali, where crop is rare, stratified by greenness (NDVI) to get "
              "enough crop points: 447 points, only 10 crop."),
             ("Label by eye, blind to all maps",
              "Monthly 3 m PlanetScope images, Sentinel-2 and Google Earth across the season. Crop = sowing, "
              "growth or harvest visible in the greenness time series. ≥ 2 labellers per point."),
             ("Score every map on the same points",
              "Accuracy, precision, recall and F1 with standard errors. 11 maps, 8 countries, no South Sudan.")]
    for i, (h, b) in enumerate(steps):
        y = (1.5, 2.6, 3.9)[i]
        circle_num(s, 0.69, y + 0.03, 0.42, str(i + 1))
        text(s, 1.22, y, 3.9, 0.35, [h], size=14, bold=True, color=TEAL)
        text(s, 1.22, y + 0.34, 3.9, 0.8, [b], size=11)

    text(s, 5.4, 1.45, 3.9, 0.3, ["Mali: 2.2% of the points are crop"], size=13, bold=True)
    rows = [("Esri", "0.97", "0.00"), ("DE Africa", "0.96", "0.46"), ("WorldCover", "0.95", "0.43"),
            ("GLAD", "0.96", "0.33"), ("ASAP", "0.95", "0.21"), ("GFSAD", "0.91", "0.15"),
            ("Mean of 11 maps", "0.92", "0.21")]
    table(s, 5.4, 1.8, 3.9, [1.9, 1.0, 1.0], rows, ("Map", "Accuracy", "F1"), size=11, head_size=11,
          red_rows=(0,))
    text(s, 5.4, 4.3, 3.9, 0.8, [
        ("Accuracy misleads when crop is rare: Esri scores 0.97 and finds no crop at all. Aweil is like Mali, "
         "so we report precision, recall and F1.", {"size": 11, "bold": True, "color": RED})], size=11)
    source(s, "Kerner et al. (2024), How accurate are existing land cover maps for agriculture in Sub-Saharan "
              "Africa? Scientific Data 11:486, Tables 1 and 4. Precision = user's accuracy, recall = producer's.")
    notes(s, "Kerner et al. 2024 is the audit we extend. Sample design: uniform random points in 7 countries; Mali "
             "stratified by four mean-annual-NDVI bands with equal points each, because cropland is so rare there. "
             "Response design: trained people looked at monthly PlanetScope composites (3 m) plus Sentinel-2 and "
             "Google Earth Pro in Collect Earth Online, month by month over the growing season; 'active cropland' "
             "= sowing, growing or harvesting seen within 12 months. Labellers never saw the maps. They discarded "
             "points where labellers disagreed. Metrics from the area-weighted error matrix (Stehman & Foody 2019; "
             "Olofsson et al. 2014). Findings: all maps agree on crop in < 0.5% of pixels; mean F1 < 0.7 in 7 of 8 "
             "countries; Mali lowest (mean F1 0.21). Esri's Mali row: accuracy 0.97 with precision, recall and F1 "
             "all 0, because it called every point non-crop.")


def slide_ro(s):
    card(s, 0.69, 1.5, 5.5, 3.6)
    text(s, 0.87, 1.6, 5.2, 0.35, ["Proposed RO"], size=13, bold=True, color=TEAL)
    text(s, 0.87, 1.95, 5.2, 0.8, [
        "Measure how much cropland ZOA's areas had in 2022, and how accurately public maps, and a readable map "
        "of our own, capture it."], size=14, bold=True)
    qs = [("How much cropland?", "hectares ± 95% CI from ≈ 450 labelled points, independent of any map"),
          ("How wrong are the public maps?", "7 maps scored on held-out points: precision, recall, F1, area"),
          ("Can readable rules do better here?", "our 2022 map, scored on the same points"),
          ("Does the choice matter for AA?", "2022 flood × each map: would a trigger have fired?")]
    for i, (q, a) in enumerate(qs):
        y = 2.85 + i * 0.55
        circle_num(s, 0.9, y + 0.04, 0.34, str(i + 1))
        text(s, 1.35, y, 4.7, 0.3, [q], size=12, bold=True)
        text(s, 1.35, y + 0.25, 4.7, 0.3, [a], size=10.5, color=MUTED)

    text(s, 6.5, 1.5, 2.8, 0.35, ["Scope"], size=13, bold=True, color=TEAL)
    scope = [("Where", "5 Aweil counties (Lol river) + Bor South: ZOA's areas"),
             ("When", "2022 season: flooding, and CFSAM figures to compare"),
             ("Map", "whole country, but checked only in ZOA's areas"),
             ("Out", "the flood forecast (documented in prediction/)"),
             ("Gap", "Kerner et al. skipped South Sudan; we found no accuracy check here")]
    for i, (k, v) in enumerate(scope):
        y = 1.9 + i * 0.64
        text(s, 6.5, y, 0.75, 0.55, [k], size=11, bold=True, color=RED)
        text(s, 7.25, y, 2.05, 0.6, [v], size=11)
    notes(s, "RO as in deliverables/CROPLAND_PLAN.html (28 Sep 2026): 'how much cropland does South Sudan have in "
             "2022, and how wrong are the public maps that ZOA's trigger would use?' The area estimate is the main "
             "result and does not depend on any map. The national map is a second result, checked only in ZOA's "
             "areas. Week 4 overlays the observed 2022 flood extent, not a forecast. Areas: Q&A 2, slide 3.")


def slide_sample(s):
    steps = [("Agreement map", "7 maps on one 10 m grid: how many say crop"),
             ("Draw ≈ 450 points", "strata 0 / 1–2 / 3+ maps say crop; most where none do"),
             ("Label blind, twice", "Sentinel-2 months + NDVI curve, Google Earth, Esri Wayback"),
             ("Split first", "150 to build the map, 300 locked away for testing"),
             ("Two results", "cropland area ± CI, and a scorecard for every map")]
    for i, (h, b) in enumerate(steps):
        x = 0.69 + i * 1.75
        circle_num(s, x + 0.56, 1.55, 0.48, str(i + 1))
        if i < 4:
            arrow(s, x + 1.1, 1.79, x + 2.25, 1.79)
        text(s, x, 2.1, 1.62, 0.35, [h], size=13, bold=True, color=TEAL, align=PP_ALIGN.CENTER)
        text(s, x, 2.45, 1.62, 0.9, [b], size=10.5, align=PP_ALIGN.CENTER)

    for j, (h, b) in enumerate([
            ("Result A: how much cropland", "Stratified estimate on all points (Olofsson et al. 2014). Valid "
             "whatever the maps say. Aweil ≈ 330 points → about ± 2.5 percentage points: enough to tell 1% from 5%."),
            ("Result B: which map to trust", "7 public maps + ours on the same 300 test points: precision, "
             "recall, F1 and area, with error bars. No map is judged on points it has seen.")]):
        x = 0.69 + j * 4.4
        card(s, x, 3.4, 4.22, 1.25)
        text(s, x + 0.15, 3.47, 3.95, 0.35, [h], size=13, bold=True, color=TEAL)
        text(s, x + 0.15, 3.8, 3.95, 0.85, [b], size=10.5)
    text(s, 0.69, 4.75, 8.62, 0.45, [
        ("Changed from Kerner: we keep points where labellers disagree (report Cohen's kappa), and label from "
         "Sentinel-2 and Google Earth because the free Planet imagery (NICFI) ended in 2025.",
         {"size": 10.5, "color": INK})], size=10.5)
    notes(s, "Why the sample comes first: rules need thresholds, thresholds need labelled points, and a map must be "
             "judged on points it never saw. Strata come from the agreement map; most points go to the 'no map says "
             "crop' land because the uncertainty of the area comes almost entirely from that big stratum. Aweil "
             "~330 points, Bor South ~120 (wider interval). Per point: crop / not crop / unsure, cues seen, imagery "
             "and date, confidence 1-3. Pilot of 60 points near Turalei first. About 1.5 min per point: 450 x 2 = "
             "~22 hours, ~4.5 hours each. Estimator: p = sum W_h p_h, SE from Olofsson et al. 2014 and Stehman 2014 "
             "(strata are agreement levels, not one map's classes).")


def slide_our_map(s):
    rows = [("Bare soil before sowing", "low NDVI, bare-soil index Apr–May"),
            ("Green-up with the rains", "NDVI rise Jul–Aug; radar if cloudy"),
            ("Harvest, grass stays green", "NDVI drop Sep–Oct vs 500 m around"),
            ("Regular plots with edges", "texture / edge density at 10 m"),
            ("Homesteads nearby", "distance to buildings"),
            ("Not crop: trees, water, burns", "dry-season NDVI, water, burn scars")]
    text(s, 0.69, 1.45, 5.0, 0.3, ["What the labeller sees → what the computer measures"], size=12, bold=True)
    table(s, 0.69, 1.8, 5.1, [2.2, 2.9], rows, ("Cue", "Feature (from Sentinel-2 / -1)"), size=10,
          head_size=11)

    text(s, 6.1, 1.45, 3.2, 0.3, ["Three versions, one test"], size=12, bold=True)
    versions = [("Hand rules", "baseline: if-then thresholds we set by looking at the 150 points"),
                ("Small decision tree", "main candidate: 3–4 levels, still prints as readable rules"),
                ("Black box", "benchmark only: random forest or Google's satellite embeddings")]
    for i, (h, b) in enumerate(versions):
        y = 1.82 + i * 0.78
        circle_num(s, 6.1, y + 0.03, 0.36, str(i + 1))
        text(s, 6.55, y, 2.75, 0.3, [h], size=12, bold=True, color=TEAL)
        text(s, 6.55, y + 0.27, 2.75, 0.5, [b], size=10)
    card(s, 6.1, 4.2, 3.2, 0.9)
    text(s, 6.22, 4.25, 3.0, 0.85, [
        ("Fit on the 150 calibration points → freeze the rules in a dated file → score on the 300 test points, "
         "next to the 7 public maps.", {"size": 10.5, "bold": True, "color": TEAL})], size=10.5)
    text(s, 0.69, 4.3, 5.1, 0.8, [
        ("Why rules: anyone (ZOA, a reviewer) can read why a pixel is crop; it reruns for any year; radar "
         "sees through the July–September clouds. If ours loses, the area estimate still stands.",
         {"size": 10.5})], size=10.5)
    notes(s, "From CROPLAND_PLAN.html (Algorithm section). Feature sources: Sentinel-2 surface reflectance with "
             "Cloud Score+; Sentinel-1 VH radar for the cloudy months; Google Open Buildings for homesteads; JRC "
             "Global Surface Water for permanent water; MODIS MCD64A1 for burn scars. In this savanna, crops and wild grass green up together, "
             "which is why the maps disagree; timing and context separate them: bare before sowing, harvested while "
             "the grass stays green, regular edges, near homesteads. Cloud Score+ masking and monthly medians; key "
             "features sit in the clearer months (Apr-May, Sep-Oct). We keep whichever of versions 1 and 2 is "
             "simpler at similar accuracy. The map runs for the whole country in Earth Engine but is only checked in "
             "ZOA's areas. What we can claim: 'better or worse than map X in ZOA's areas in 2022, on our test "
             "points', with error bars; not 'a better cropland map' in general.")


def slide_questions(s):
    qs = [("Is this useful to ZOA?",
           "It does not prescribe an AA action. It gives ZOA a checked cropland number and which map to trust, "
           "to process and share with communities. Is that a strong enough contribution?"),
          ("How strong is the method?",
           "What would you add? Do you already see a flaw or a gap in the sample, the labels or the test?"),
          ("Is labelling by hand worth it?",
           "≈ 450 points × 2 labellers ≈ 22 hours (≈ 4.5 h each). Good use of our time, or is there a shortcut?"),
          ("Is 2022 a useful year?",
           "Many maps stop at 2019 (GLAD, DE Africa). Is a 2022 check useful to ZOA, or would a more recent "
           "year matter more?")]
    for i, (q, why) in enumerate(qs):
        x = 0.69 + (i % 2) * 4.4
        y = 1.55 + (i // 2) * 1.8
        card(s, x, y, 4.22, 1.6)
        circle_num(s, x + 0.15, y + 0.17, 0.42, str(i + 1))
        text(s, x + 0.7, y + 0.15, 3.4, 0.45, [q], size=15, bold=True, color=TEAL, anchor=MSO_ANCHOR.MIDDLE)
        text(s, x + 0.7, y + 0.62, 3.4, 0.95, [why], size=11.5)
    notes(s, "Q1: the output is understanding (area with error bars, which maps miss fields), not a trigger rule. "
             "Q3: labels are the ground truth everything is scored against; LLMs are never used to label. Q4: 2022 "
             "had flooding and CFSAM figures; the pipeline can rerun for 2024 with a small new check sample "
             "(~100 points).")


def main() -> None:
    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs = Presentation(str(TEMPLATE))
    slides = list(prs.slides)
    tpl_title = copy.deepcopy(slides[3].shapes.title._element.txBody)
    layout = next(l for l in prs.slide_layouts if l.name == "Title and Content")

    # keep only the title slide of the template
    sld = prs.slides._sldIdLst
    for e in list(sld)[1:]:
        prs.part.drop_rel(e.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"))
        sld.remove(e)
    slide_title(slides[0])

    for l1, l2, fn in [("The problem: ", "a trigger that needs a cropland map", slide_why),
                       ("The story: ", "seven maps, seven answers", slide_seven_maps),
                       ("The story: ", "the maps barely agree, and they are old", slide_disagree),
                       ("Prior work: ", "how Kerner et al. (2024) audited 11 maps", slide_kerner),
                       ("Research objective: ", "what we want to confirm today", slide_ro),
                       ("Our method: ", "one labelled sample, two uses", slide_sample),
                       ("Our method: ", "our own map, tested on held-out points", slide_our_map),
                       ("Your feedback: ", "four questions", slide_questions)]:
        s = prs.slides.add_slide(layout)
        sd.drop_body(s)
        sd.set_title(s, tpl_title, l1, l2)
        fn(s)
    prs.save(str(OUT))
    print(f"wrote {OUT} ({len(prs.slides)} slides)")


if __name__ == "__main__":
    main()
